"""账本筛选导出（Excel / .xlsx，M4 收尾）。

与列表共用 `build_record_filters`：**看到什么就导出什么**（类型/类别/关键字/状态/日期/精确 id 全部生效），
不受分页限制。文件在内存里生成后一次性回传（xlsx 本质是 zip，openpyxl 走内存流，无需临时文件）。

表结构：
- 标题行：班级名 + 导出时间 + 当前筛选条件（人看不懂「无筛选」时也知道导出了什么）
- 明细：发生时间 / 方向 / 类别 / 摘要 / 金额 / 支付方式 / 状态 / 备注 / 录入人
- 合计行：收入合计、支出合计、净额（作废流水不计入合计）
"""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import ClsClass, FeeRecord
from ..security import LoginUser
from .record_service import build_record_filters, to_vo_list

HEADERS = [
    ("发生时间", 20),
    ("方向", 8),
    ("类别", 14),
    ("摘要", 30),
    ("金额（元）", 14),
    ("支付方式", 10),
    ("状态", 8),
    ("备注", 24),
    ("录入人", 12),
]

CHANNEL_TEXT = {"CASH": "现金", "WECHAT": "微信", "ALIPAY": "支付宝", "BANK": "银行卡", "OTHER": "其他"}

_HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
_TITLE_FONT = Font(bold=True, size=14)
_HEAD_FONT = Font(bold=True)
_SUM_FONT = Font(bold=True)


def _filter_text(
    type_: str | None,
    category_id: int | None,
    status: str | None,
    start: str | None,
    end: str | None,
    keyword: str | None = None,
    record_id: int | None = None,
    channel: str | None = None,
) -> str:
    """表头里回显「导出了什么」。与 `build_record_filters` 同口径：`all` / 空 = 不筛选。"""
    parts: list[str] = []
    t = (type_ or "").strip().upper()
    if t and t != "ALL":
        parts.append("方向=" + ("收入" if t == "INCOME" else "支出"))
    if category_id is not None:
        parts.append(f"类别 #{category_id}")
    s = (status or "").strip().upper()
    if s and s != "ALL":
        parts.append("状态=" + ("正常" if s == "NORMAL" else "作废"))
    ch = (channel or "").strip().upper()
    if ch and ch != "ALL":
        parts.append("渠道=" + CHANNEL_TEXT.get(ch, ch))
    if start or end:
        parts.append(f"日期={start or '不限'} ~ {end or '不限'}")
    kw = keyword.strip() if keyword else ""
    if kw:
        parts.append(f"关键词「{kw}」")
    if record_id is not None:
        parts.append(f"指定流水 #{record_id}")
    return "、".join(parts) if parts else "全部流水"


def build_export_bytes(
    db: Session,
    user: LoginUser,
    *,
    type_: str | None,
    category_id: int | None,
    keyword: str | None,
    status: str | None,
    record_id: int | None,
    start: str | None,
    end: str | None,
    channel: str | None = None,
) -> bytes:
    where = build_record_filters(
        user,
        type_=type_,
        category_id=category_id,
        keyword=keyword,
        status=status,
        record_id=record_id,
        start=start,
        end=end,
        channel=channel,
    )
    rows = list(
        db.scalars(
            select(FeeRecord)
            .where(where)
            .order_by(FeeRecord.occurred_at.asc(), FeeRecord.id.asc())
        )
    )
    records = to_vo_list(db, rows)

    clazz = db.get(ClsClass, user.class_id)
    class_name = clazz.class_name if clazz is not None else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "账本明细"

    ws.cell(row=1, column=1, value=f"{class_name} 班费账本明细").font = _TITLE_FONT
    ws.cell(
        row=2,
        column=1,
        value=f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        f"　共 {len(records)} 条　筛选：{_filter_text(type_, category_id, status, start, end, keyword, record_id, channel)}",
    ).font = Font(size=10, color="5A6779")

    header_row = 4
    for idx, (title, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=title)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = width

    income = Decimal("0.00")
    expense = Decimal("0.00")
    row_no = header_row
    for rec in records:
        row_no += 1
        amount = Decimal(str(rec["amount"]))
        # 作废流水照常列出（可核对），但不计入合计
        if rec["status"] == "NORMAL":
            if rec["type"] == "INCOME":
                income += amount
            else:
                expense += amount
        cells = [
            rec["occurredAt"],
            "收入" if rec["type"] == "INCOME" else "支出",
            rec["categoryName"] or "",
            rec["title"],
            float(amount),
            CHANNEL_TEXT.get(rec["channel"] or "", rec["channel"] or ""),
            "正常" if rec["status"] == "NORMAL" else "已作废",
            rec["remark"] or "",
            rec["createdByName"] or "",
        ]
        for idx, value in enumerate(cells, start=1):
            cell = ws.cell(row=row_no, column=idx, value=value)
            if idx == 5:  # 金额右对齐 + 两位小数
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    row_no += 1
    ws.cell(row=row_no, column=4, value="合计").font = _SUM_FONT
    ws.cell(row=row_no, column=5, value=f"收入 {income:,.2f}　支出 {expense:,.2f}　净额 {income - expense:,.2f}").font = _SUM_FONT

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:I{row_no - 1}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename() -> str:
    return f"班费账本_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
