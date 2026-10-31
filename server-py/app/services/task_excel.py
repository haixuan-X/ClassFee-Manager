"""班级任务的 Excel 导出（第二十五轮）。

只导**明细**（谁填了什么、谁还没填），因为任务列表页给所有人看的都只是汇总，
明细是发布人/班主任才有权看的那一份——导出的口径必须和 `/detail` 一致，
否则就成了绕过 403 的越权口子。

导出前先过一次 `task_service.get_detail`：权限判定只有那一处，
这里不再重复实现「谁能看」的规则。
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..errors import BizException
from ..security import LoginUser
from . import task_service as svc

_TITLE_FONT = Font(bold=True, size=14)
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="4A7EBB")
_MUTED_FONT = Font(size=10, color="666666")

ROLE_TEXT = {
    "DEAN": "院长",
    "TEACHER": "班主任",
    "MONITOR": "班长",
    "ADMIN": "生活委员",
    "MEMBER": "普通成员",
    "GUEST": "访客",
}


def export_detail(
    db: Session,
    user: LoginUser,
    task_id: int,
    only_ids: list[int] | None = None,
    label: str | None = None,
) -> tuple[bytes, str]:
    """导出某个任务的回答。权限由 get_detail 把关（无权会抛 403）。

    **班级任务**一行 = 一个学生账号；**学院任务**一行 = 一个学生，
    多一列「班级」（按班分组排好序）。两套口径由 `get_detail` 分流。

    `only_ids` = **按状态筛选的导出**（第四十九轮）：只导这些行，
    没有匹配学生的班**整个不出现**（留着一排空表头会让人以为漏了人）。
    `label` 只是这份筛选的名字，进副标题和文件名，不参与过滤。

    过滤在这里、语义在前端：「哪几个 recordId 算未返校」由 `taskStatus.ts`
    算好传上来，这里不重复判定——两处各写一份关键词迟早漂。
    """
    detail = svc.get_detail(db, user, task_id)
    is_college = detail.get("scope") == "COLLEGE"
    classes = detail.get("classes") or []

    #: 筛选态：前端传的是**屏幕上那些行**的 id。学院任务的行键是 recordId，
    #: 班级任务的行键是 userId（两套明细结构不同），按各自有的那个键匹配。
    only = set(only_ids) if only_ids else None

    def _keep(r: dict) -> bool:
        if only is None:
            return True
        key = r.get("recordId") if is_college else r.get("userId")
        return key in only

    wb = Workbook()
    ws = wb.active
    ws.title = "任务回答"

    ws.cell(row=1, column=1, value=f"{detail['title']} · 导出 {datetime.now():%Y-%m-%d %H:%M}").font = _TITLE_FONT
    # 筛选态必须写在**第一屏**（第 2 行）：否则拿到文件的人会以为这是全量，
    # 拿着 6 行的名单去对 35 个人的班，对不上还以为是漏导
    if only is None:
        kept = 0
    elif is_college:
        kept = sum(1 for c in classes for r in c["rows"] if _keep(r))
    else:
        kept = sum(1 for r in detail["rows"] if _keep(r))
    if only is not None and kept == 0:
        # 一个 id 都没命中 = 传进来的行已经不在明细里（改过/刷过）。
        # 发一份只有表头的空文件比报错更糟：拿文件的人会以为「这个班没人」
        raise BizException.bad_request("这些行不在明细里了，打开详情重新筛选后再导出")
    ws.cell(
        row=2,
        column=1,
        value=(
            f"题目：{detail['question']}（{detail['qtypeText']}）　"
            f"已填 {detail['answered']} / 共 {detail['unanswered'] + detail['answered']} 人"
            + (f"　带说明 {detail['noteCount']} 条" if detail["allowText"] else "")
            + (f"　筛选：{label or '所选行'}（{kept} 人）" if only is not None else "")
        ),
    ).font = _MUTED_FONT

    if is_college:
        headers = ["班级", "班主任", "学号", "姓名", "是否已填", "状态", "补充说明", "提交时间"]
        col_widths = [16, 12, 14, 12, 10, 16, 28, 18]
    else:
        headers = ["姓名", "登录名", "角色", "是否已填", "答案", "补充说明", "提交时间"]
        col_widths = [16, 16, 12, 10, 28, 28, 18]

    # ⚠ 循环变量不能叫 `label`：那会把**函数参数** `label`（筛选名）覆盖掉，
    # 文件名就会变成「…-提交时间-回答.xlsx」（表头最后一个值）。e2e 抓到过一次。
    for i, col_title in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=col_title)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = col_widths[i - 1]

    if is_college:
        r_i = 4
        for cls in classes:
            rows = [r for r in cls["rows"] if _keep(r)]
            if only is not None and not rows:
                continue  # 这个班没有命中筛选的人 → 整段不出现在文件里
            # 班小标题：一眼看出「这个班交了多少」
            # 筛选态下原来的「已填 x / 共 y 人」是**全量**口径，配几张筛选后的行
            # 会读成「这个班就这几个人」——换成筛选口径的人数
            head = ws.cell(
                row=r_i,
                column=1,
                value=(
                    (
                        f"{cls['className']}　筛选出 {len(rows)} 人"
                        if only is not None
                        else f"{cls['className']}　已填 {cls['answered']} / 共 {cls['total']} 人"
                    )
                    + (f"　班主任 {'、'.join(cls['teachers'])}" if cls["teachers"] else "")
                ),
            )
            head.font = Font(bold=True, color="1F4E79")
            head.fill = PatternFill("solid", fgColor="DDEBF7")
            r_i += 1
            # 未填的排前面（和明细页同序）
            rows = sorted(rows, key=lambda r: (r["answered"], r["name"]))
            for r in rows:
                ws.cell(row=r_i, column=1, value=cls["className"])
                ws.cell(row=r_i, column=2, value="、".join(cls["teachers"]))
                ws.cell(row=r_i, column=3, value=r["studentNo"])
                ws.cell(row=r_i, column=4, value=r["name"])
                ws.cell(row=r_i, column=5, value="是" if r["answered"] else "未填")
                ws.cell(row=r_i, column=6, value=r["value"] or "")
                ws.cell(row=r_i, column=7, value=r["note"] or "")
                ws.cell(row=r_i, column=8, value=r["submittedAt"] or "")
                r_i += 1
    else:
        rows = [r for r in detail["rows"] if _keep(r)]
        for r_i, r in enumerate(rows, start=4):
            # 未填的人排在前面（service 已排好序），导出保持同一顺序，便于对照催交
            ws.cell(row=r_i, column=1, value=r["name"])
            ws.cell(row=r_i, column=2, value=r["username"])
            ws.cell(row=r_i, column=3, value=ROLE_TEXT.get(r["role"], r["role"]))
            ws.cell(row=r_i, column=4, value="是" if r["answered"] else "未填")
            ws.cell(row=r_i, column=5, value=r["value"] or "")
            ws.cell(row=r_i, column=6, value=r["note"] or "")
            ws.cell(row=r_i, column=7, value=r["submittedAt"] or "")

    buf = BytesIO()
    wb.save(buf)
    # 筛选态进文件名：下载栏里三个文件并排时，才知道哪个是「未返校」那份
    suffix = f"-{label}" if only is not None and label else ""
    return buf.getvalue(), f"{detail['title']}{suffix}-回答.xlsx"


def export_rollcall(
    db: Session, user: LoginUser, task_id: int, class_id: int | None = None
) -> tuple[bytes, str]:
    """班主任导出**自己班**的点名结果（第四十五轮）。

    和 `export_detail` 是两个端点：那个是院长看全院明细（`can_see_detail` 拦），
    这个是班主任导自己填的那份（`get_rollcall` 里按 `managed_class_ids` 裁）。
    """
    detail = svc.get_rollcall(db, user, task_id, class_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "点名结果"

    ws.cell(row=1, column=1, value=f"{detail['title']} · {detail['className']} · 导出 {datetime.now():%Y-%m-%d %H:%M}").font = _TITLE_FONT
    ws.cell(
        row=2,
        column=1,
        value=(
            f"题目：{detail['question']}（{detail['qtypeText']}）　"
            f"已填 {detail['answered']} / 共 {detail['total']} 人"
            + (f"　带说明 {detail['noteCount']} 条" if detail["allowText"] else "")
        ),
    ).font = _MUTED_FONT

    if detail["allowText"]:
        headers = ["学号", "姓名", "状态", "补充说明", "提交时间"]
        col_widths = [14, 12, 16, 28, 18]
    else:
        headers = ["学号", "姓名", "状态", "提交时间"]
        col_widths = [14, 12, 16, 18]

    for i, label in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=label)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = col_widths[i - 1]

    for r_i, r in enumerate(detail["rows"], start=4):
        ws.cell(row=r_i, column=1, value=r["studentNo"])
        ws.cell(row=r_i, column=2, value=r["name"])
        ws.cell(row=r_i, column=3, value=r["value"] or "未填")
        if detail["allowText"]:
            ws.cell(row=r_i, column=4, value=r["note"] or "")
            ws.cell(row=r_i, column=5, value=r["submittedAt"] or "")
        else:
            ws.cell(row=r_i, column=4, value=r["submittedAt"] or "")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{detail['title']}-{detail['className']}-点名结果.xlsx"


def export_overview(db: Session, user: LoginUser) -> tuple[bytes, str]:
    """导出**任务总览**：所有任务一览 + 回收情况（仅院长可调）。

    和「单条任务导出」不同：这里是给院长看全局的一页纸——每条任务的
    标题、范围、发布人、开放状态、回收进度、截止、发布时间。
    """
    from sqlalchemy import select

    from ..models import ClsTask, SysUser

    tasks = list(
        db.execute(
            select(ClsTask).order_by(ClsTask.created_at.desc(), ClsTask.id.desc())
        ).scalars()
    )
    rows = svc._build_rows(db, user, tasks)

    publisher: dict[int, str] = {}
    if tasks:
        for uid, rn, un in db.execute(
            select(SysUser.id, SysUser.real_name, SysUser.username).where(
                SysUser.id.in_([t.created_by for t in tasks])
            )
        ):
            publisher[int(uid)] = rn or un or ""

    wb = Workbook()
    ws = wb.active
    ws.title = "任务总览"

    ws.cell(
        row=1, column=1, value=f"任务总览 · 导出 {datetime.now():%Y-%m-%d %H:%M}"
    ).font = _TITLE_FONT
    ws.cell(row=2, column=1, value=f"共 {len(rows)} 条任务").font = _MUTED_FONT

    headers = ["标题", "范围", "发布人", "状态", "已填", "应答总数", "回收率", "截止", "发布时间", "置顶"]
    for i, label in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=label)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = 24 if i == 1 else 14

    for r_i, r in enumerate(rows, start=4):
        s = r.get("summary") or {}
        answered = int(s.get("answered") or 0)
        total = int(s.get("total") or 0)
        rate = f"{round(answered * 100 / total)}%" if total else "—"
        scope = (
            "学院任务"
            if r.get("scope") == "COLLEGE"
            else (r.get("className") or "班级任务")
        )
        ws.cell(row=r_i, column=1, value=r.get("title") or "")
        ws.cell(row=r_i, column=2, value=scope)
        ws.cell(row=r_i, column=3, value=publisher.get(int(r.get("createdBy") or 0), ""))
        ws.cell(row=r_i, column=4, value="开放中" if r.get("open") else "已关闭")
        ws.cell(row=r_i, column=5, value=answered)
        ws.cell(row=r_i, column=6, value=total)
        ws.cell(row=r_i, column=7, value=rate)
        ws.cell(row=r_i, column=8, value=r.get("deadline") or "不限")
        ws.cell(row=r_i, column=9, value=r.get("createdAt") or "")
        ws.cell(row=r_i, column=10, value="是" if r.get("pinned") else "")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), "任务总览.xlsx"


def xlsx_response(pair: tuple[bytes, str]):
    """包成带附件头的 HTTP 响应。

    中文文件名必须走 RFC 5987 的 `filename*`，否则浏览器拿到乱码——
    与 students 的导出同一套处理。
    """
    from fastapi.responses import Response

    content, filename = pair
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="tasks.xlsx"; filename*=UTF-8\'\'{quote(filename)}'
            )
        },
    )